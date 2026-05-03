import csv
import os
import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font, Alignment, PatternFill
except ImportError:
    print("Instalando openpyxl...")
    install('openpyxl')
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font, Alignment, PatternFill

output_dir = r"C:\Users\ASUS\.gemini\antigravity\scratch"

def run():
    # 1. Read the CSVs
    nodo_data = []
    with open(os.path.join(output_dir, "Prueba_BD_NODO.csv"), "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader: nodo_data.append(row)
        
    pext_data = []
    with open(os.path.join(output_dir, "Prueba_BD_PEXT.csv"), "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader: pext_data.append(row)

    iao_data = []
    with open(os.path.join(output_dir, "Prueba_BD_IAO.csv"), "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader: iao_data.append(row)

    # 2. Process data for Dashboard
    # Metric 1: Total Metraje FO
    total_fo = sum(float(r["Metraje_Promedio_FO"]) for r in pext_data if r["Metraje_Promedio_FO"].replace('.', '', 1).isdigit())
    
    # Metric 2: Total IAOs por Tipo
    iaos_por_tipo = {}
    for r in iao_data:
        tipo = r["Tipo_IAO"]
        iaos_por_tipo[tipo] = iaos_por_tipo.get(tipo, 0) + 1
        
    # Metric 3: Total IAOs por Cuadrilla
    iaos_por_cuadrilla = {}
    for r in iao_data:
        c = r["Cuadrilla_IAO"]
        iaos_por_cuadrilla[c] = iaos_por_cuadrilla.get(c, 0) + 1

    # Metric 4: Total Postes Nuevos vs Electricos
    postes_nuevos = sum(int(r["Cant_Postes_Nuevos"]) for r in pext_data if r["Cant_Postes_Nuevos"].isdigit())
    postes_electricos = sum(int(r["Cant_Postes_Electricos"]) for r in pext_data if r["Cant_Postes_Electricos"].isdigit())
    postes_terceros = sum(int(r["Cant_Postes_Terceros"]) for r in pext_data if r["Cant_Postes_Terceros"].isdigit())

    # 3. Create Workbook
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = "DASHBOARD"
    ws_datos = wb.create_sheet("Datos_Calculados")
    
    # Write aggregated data to Datos_Calculados
    ws_datos.append(["Métrica", "Valor"])
    ws_datos.append(["Total FO (Metros)", total_fo])
    
    ws_datos.append([])
    ws_datos.append(["Tipo IAO", "Cantidad"])
    row_iao_tipo_start = ws_datos.max_row + 1
    for k, v in iaos_por_tipo.items():
        ws_datos.append([k, v])
    row_iao_tipo_end = ws_datos.max_row

    ws_datos.append([])
    ws_datos.append(["Cuadrilla", "IAOs Terminadas"])
    row_cuad_start = ws_datos.max_row + 1
    for k, v in iaos_por_cuadrilla.items():
        ws_datos.append([k, v])
    row_cuad_end = ws_datos.max_row
    
    ws_datos.append([])
    ws_datos.append(["Tipo Poste", "Cantidad"])
    row_poste_start = ws_datos.max_row + 1
    ws_datos.append(["Nuevos Pronatel", postes_nuevos])
    ws_datos.append(["Eléctricos MT/BT", postes_electricos])
    ws_datos.append(["Terceros", postes_terceros])
    row_poste_end = ws_datos.max_row

    # 4. Format DASHBOARD sheet
    ws_dash.sheet_properties.tabColor = "1072BA"
    ws_dash.column_dimensions['A'].width = 2
    ws_dash.column_dimensions['B'].width = 40
    ws_dash.column_dimensions['C'].width = 20
    ws_dash.column_dimensions['J'].width = 15
    
    # Title
    ws_dash["B2"] = "DASHBOARD DIRECTIVO - PROYECTO FTTH PRONATEL"
    ws_dash["B2"].font = Font(size=20, bold=True, color="FFFFFF")
    ws_dash["B2"].fill = PatternFill(start_color="1072BA", end_color="1072BA", fill_type="solid")
    ws_dash["C2"].fill = PatternFill(start_color="1072BA", end_color="1072BA", fill_type="solid")

    ws_dash["B4"] = "KPIs Principales del Proyecto"
    ws_dash["B4"].font = Font(size=14, bold=True, color="1072BA")
    
    ws_dash["B5"] = "Total de FO Desplegada (Metros):"
    ws_dash["C5"] = total_fo
    ws_dash["C5"].number_format = '#,##0.00'
    ws_dash["C5"].font = Font(bold=True)
    
    ws_dash["B6"] = "Total de IAOs Instaladas:"
    ws_dash["C6"] = len(iao_data)
    ws_dash["C6"].font = Font(bold=True)
    
    ws_dash["B7"] = "Postes Eléctricos Intervenidos:"
    ws_dash["C7"] = postes_electricos
    ws_dash["C7"].font = Font(bold=True)

    # Chart 1: IAOs por Tipo (Pie Chart)
    pie = PieChart()
    labels = Reference(ws_datos, min_col=1, min_row=row_iao_tipo_start, max_row=row_iao_tipo_end)
    data = Reference(ws_datos, min_col=2, min_row=row_iao_tipo_start-1, max_row=row_iao_tipo_end)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.title = "Distribución de IAOs por Tipo"
    ws_dash.add_chart(pie, "B9")

    # Chart 2: IAOs por Cuadrilla (Bar Chart)
    bar1 = BarChart()
    labels1 = Reference(ws_datos, min_col=1, min_row=row_cuad_start, max_row=row_cuad_end)
    data1 = Reference(ws_datos, min_col=2, min_row=row_cuad_start-1, max_row=row_cuad_end)
    bar1.add_data(data1, titles_from_data=True)
    bar1.set_categories(labels1)
    bar1.title = "Productividad: IAOs por Cuadrilla"
    bar1.legend = None
    bar1.y_axis.title = "Cantidad"
    ws_dash.add_chart(bar1, "J9")
    
    # Chart 3: Postes (Bar Chart)
    bar2 = BarChart()
    labels2 = Reference(ws_datos, min_col=1, min_row=row_poste_start, max_row=row_poste_end)
    data2 = Reference(ws_datos, min_col=2, min_row=row_poste_start-1, max_row=row_poste_end)
    bar2.add_data(data2, titles_from_data=True)
    bar2.set_categories(labels2)
    bar2.title = "Infraestructura: Uso de Postes"
    bar2.legend = None
    ws_dash.add_chart(bar2, "B24")

    # Add Raw Data sheets
    ws_n = wb.create_sheet("BD_NODO")
    ws_n.append(list(nodo_data[0].keys()))
    for r in nodo_data: ws_n.append(list(r.values()))
    
    ws_p = wb.create_sheet("BD_PEXT")
    ws_p.append(list(pext_data[0].keys()))
    for r in pext_data: ws_p.append(list(r.values()))

    ws_i = wb.create_sheet("BD_IAO")
    ws_i.append(list(iao_data[0].keys()))
    for r in iao_data: ws_i.append(list(r.values()))

    # Hide Datos_Calculados
    ws_datos.sheet_state = 'hidden'

    out_path = os.path.join(output_dir, "Dashboard_Completo_FTTH.xlsx")
    wb.save(out_path)
    print(f"Dashboard generado con éxito en: {out_path}")

if __name__ == "__main__":
    run()
